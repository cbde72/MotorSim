from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import hashlib
import os
import re
import tempfile

import numpy as np


GENERATOR_MAP_CACHE_FORMAT_VERSION = 2


@dataclass(frozen=True, slots=True)
class GeneratorMapBranch:
    angle_deg: np.ndarray
    torque_Nm: np.ndarray
    voltage_V: np.ndarray
    electrical_power_W: np.ndarray


@dataclass(frozen=True, slots=True)
class GeneratorMapOperatingPoint:
    increasing: GeneratorMapBranch
    decreasing: GeneratorMapBranch


@dataclass(frozen=True, slots=True)
class GeneratorTorqueMap:
    source_path: Path
    resistances_ohm: np.ndarray
    operating_points: tuple[GeneratorMapOperatingPoint, ...]
    no_load: GeneratorMapOperatingPoint
    load_resistance_ohm: float
    skalierung_faktor: float
    angle_at_x_min_deg: float
    angle_at_x_max_deg: float
    include_no_load_torque: bool
    angle_out_of_range: str
    resistance_out_of_range: str
    max_abs_torque_Nm: float
    max_electrical_power_W: float


@dataclass(frozen=True, slots=True)
class GeneratorMapValue:
    angle_deg: float
    torque_Nm: float
    no_load_torque_Nm: float
    load_torque_Nm: float
    voltage_V: float
    electrical_power_W: float
    out_of_range: bool


def _numeric_rows(ws) -> list[tuple[float, float, float, float, float]]:
    rows: list[tuple[float, float, float, float, float]] = []
    for row in range(3, int(ws.max_row) + 1):
        values = [ws.cell(row, col).value for col in (2, 3, 4, 5, 6)]
        if not all(isinstance(value, (int, float)) for value in values[:3]):
            continue
        rows.append((float(values[0]), float(values[1]), float(values[2]), float(values[3] or 0.0), float(values[4] or 0.0)))
    if len(rows) < 3:
        raise ValueError(f"generator map sheet {ws.title!r} contains fewer than three numeric rows")
    return rows


def _branch(rows: list[tuple[float, float, float, float, float]], increasing: bool) -> GeneratorMapBranch:
    selected: list[tuple[float, float, float, float]] = []
    for index, row in enumerate(rows):
        if index == 0:
            slope = rows[1][1] - row[1]
        elif index == len(rows) - 1:
            slope = row[1] - rows[index - 1][1]
        else:
            slope = rows[index + 1][1] - rows[index - 1][1]
        if (slope >= 0.0) == increasing:
            selected.append((row[1], row[2], row[3], row[4]))
    selected.sort(key=lambda item: item[0])
    deduplicated: list[tuple[float, float, float, float]] = []
    for item in selected:
        if deduplicated and abs(item[0] - deduplicated[-1][0]) < 1.0e-9:
            deduplicated[-1] = tuple(0.5 * (a + b) for a, b in zip(deduplicated[-1], item))
        else:
            deduplicated.append(item)
    if len(deduplicated) < 2:
        raise ValueError("generator map cannot form both direction-dependent angle branches")
    values = np.asarray(deduplicated, dtype=np.float64)
    return GeneratorMapBranch(values[:, 0], values[:, 1], values[:, 2], values[:, 3])


def _operating_point(ws) -> GeneratorMapOperatingPoint:
    rows = _numeric_rows(ws)
    return GeneratorMapOperatingPoint(increasing=_branch(rows, True), decreasing=_branch(rows, False))


def _scale_cache_token(skalierung_faktor: float) -> str:
    return format(float(skalierung_faktor), ".12g").replace("-", "m").replace(".", "p").replace("+", "")


def generator_map_cache_path(source_path: str | Path, skalierung_faktor: float = 1.0) -> Path:
    source = Path(source_path)
    return source.with_suffix(f".generator_map.scale_{_scale_cache_token(skalierung_faktor)}.npz")


def _source_sha256(source: Path) -> str:
    digest = hashlib.sha256()
    with source.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _branch_cache_values(prefix: str, branch: GeneratorMapBranch) -> dict[str, np.ndarray]:
    return {
        f"{prefix}_angle_deg": np.asarray(branch.angle_deg, dtype=np.float64),
        f"{prefix}_torque_Nm": np.asarray(branch.torque_Nm, dtype=np.float64),
        f"{prefix}_voltage_V": np.asarray(branch.voltage_V, dtype=np.float64),
        f"{prefix}_electrical_power_W": np.asarray(branch.electrical_power_W, dtype=np.float64),
    }


def _point_cache_values(prefix: str, point: GeneratorMapOperatingPoint) -> dict[str, np.ndarray]:
    values = _branch_cache_values(f"{prefix}_increasing", point.increasing)
    values.update(_branch_cache_values(f"{prefix}_decreasing", point.decreasing))
    return values


def _branch_from_cache(data: np.lib.npyio.NpzFile, prefix: str) -> GeneratorMapBranch:
    return GeneratorMapBranch(
        angle_deg=np.asarray(data[f"{prefix}_angle_deg"], dtype=np.float64),
        torque_Nm=np.asarray(data[f"{prefix}_torque_Nm"], dtype=np.float64),
        voltage_V=np.asarray(data[f"{prefix}_voltage_V"], dtype=np.float64),
        electrical_power_W=np.asarray(data[f"{prefix}_electrical_power_W"], dtype=np.float64),
    )


def _point_from_cache(data: np.lib.npyio.NpzFile, prefix: str) -> GeneratorMapOperatingPoint:
    return GeneratorMapOperatingPoint(
        increasing=_branch_from_cache(data, f"{prefix}_increasing"),
        decreasing=_branch_from_cache(data, f"{prefix}_decreasing"),
    )


def _write_cache(cache_path: Path, source: Path, source_hash: str, skalierung_faktor: float, resistances: np.ndarray, points: tuple[GeneratorMapOperatingPoint, ...], no_load: GeneratorMapOperatingPoint) -> None:
    values: dict[str, np.ndarray] = {
        "format_version": np.asarray(GENERATOR_MAP_CACHE_FORMAT_VERSION, dtype=np.int64),
        "source_filename": np.asarray(source.name),
        "source_sha256": np.asarray(source_hash),
        "source_mtime_ns": np.asarray(source.stat().st_mtime_ns, dtype=np.int64),
        "skalierung_faktor": np.asarray(skalierung_faktor, dtype=np.float64),
        "resistances_ohm": np.asarray(resistances, dtype=np.float64),
    }
    values.update(_point_cache_values("no_load", no_load))
    for index, point in enumerate(points):
        values.update(_point_cache_values(f"load_{index}", point))
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(prefix=f".{cache_path.stem}.", suffix=".npz", dir=cache_path.parent, delete=False) as stream:
            temp_path = Path(stream.name)
        np.savez_compressed(temp_path, **values)
        os.replace(temp_path, cache_path)
    finally:
        if temp_path is not None and temp_path.exists():
            temp_path.unlink()


def _read_cache(cache_path: Path, expected_hash: str | None, expected_scale: float = 1.0) -> tuple[np.ndarray, tuple[GeneratorMapOperatingPoint, ...], GeneratorMapOperatingPoint] | None:
    if not cache_path.exists():
        return None
    try:
        with np.load(cache_path, allow_pickle=False) as data:
            if int(np.asarray(data["format_version"]).item()) != GENERATOR_MAP_CACHE_FORMAT_VERSION:
                return None
            cached_hash = str(np.asarray(data["source_sha256"]).item())
            if expected_hash is not None and cached_hash != expected_hash:
                return None
            if not np.isclose(float(np.asarray(data["skalierung_faktor"]).item()), expected_scale, rtol=0.0, atol=1.0e-12):
                return None
            resistances = np.asarray(data["resistances_ohm"], dtype=np.float64).copy()
            points = tuple(_point_from_cache(data, f"load_{index}") for index in range(int(resistances.size)))
            no_load = _point_from_cache(data, "no_load")
            return resistances, points, no_load
    except (KeyError, OSError, ValueError):
        return None


def _scaled_branch(branch: GeneratorMapBranch, factor: float) -> GeneratorMapBranch:
    return GeneratorMapBranch(
        angle_deg=branch.angle_deg,
        torque_Nm=branch.torque_Nm * factor,
        voltage_V=branch.voltage_V,
        electrical_power_W=branch.electrical_power_W * factor,
    )


def _scaled_point(point: GeneratorMapOperatingPoint, factor: float) -> GeneratorMapOperatingPoint:
    return GeneratorMapOperatingPoint(
        increasing=_scaled_branch(point.increasing, factor),
        decreasing=_scaled_branch(point.decreasing, factor),
    )


def _load_or_create_cached_maps(source: Path, skalierung_faktor: float) -> tuple[np.ndarray, tuple[GeneratorMapOperatingPoint, ...], GeneratorMapOperatingPoint]:
    cache_path = generator_map_cache_path(source, skalierung_faktor)
    source_hash = _source_sha256(source) if source.exists() else None
    cached = _read_cache(cache_path, source_hash, skalierung_faktor)
    if cached is not None:
        return cached
    if not source.exists():
        raise ValueError(f"generator torque-map source and cache do not exist: {source}, {cache_path}")
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depends on deployment environment
        raise RuntimeError("generator_torque_map requires the optional package 'openpyxl'") from exc

    workbook = openpyxl.load_workbook(source, data_only=True, read_only=False)
    if "No Load" not in workbook.sheetnames:
        raise ValueError(f"generator torque-map workbook {source} is missing sheet 'No Load'")
    resistance_points: list[tuple[float, GeneratorMapOperatingPoint]] = []
    for sheet_name in workbook.sheetnames:
        match = re.fullmatch(r"\s*([0-9]+(?:\.[0-9]+)?)\s*Ohm\s*", sheet_name, flags=re.IGNORECASE)
        if match:
            resistance_points.append((float(match.group(1)), _operating_point(workbook[sheet_name])))
    resistance_points.sort(key=lambda item: item[0])
    if not resistance_points:
        raise ValueError(f"generator torque-map workbook {source} contains no '<value> Ohm' sheets")
    resistances = np.asarray([item[0] for item in resistance_points], dtype=np.float64)
    points = tuple(_scaled_point(item[1], skalierung_faktor) for item in resistance_points)
    no_load = _scaled_point(_operating_point(workbook["No Load"]), skalierung_faktor)
    workbook.close()
    _write_cache(cache_path, source, str(source_hash), skalierung_faktor, resistances, points, no_load)
    return resistances, points, no_load


def load_generator_torque_map(config, config_path: str | Path) -> GeneratorTorqueMap:
    source = Path(config.file)
    if not source.is_absolute():
        source = Path(config_path).resolve().parent / source
    source = source.resolve()
    skalierung_faktor = float(getattr(config, "skalierung_faktor", 1.0))
    resistances, points, no_load = _load_or_create_cached_maps(source, skalierung_faktor)
    return GeneratorTorqueMap(
        source_path=source,
        resistances_ohm=resistances,
        operating_points=points,
        no_load=no_load,
        load_resistance_ohm=float(config.load_resistance_ohm),
        skalierung_faktor=skalierung_faktor,
        angle_at_x_min_deg=float(config.angle_at_x_min_deg),
        angle_at_x_max_deg=float(config.angle_at_x_max_deg),
        include_no_load_torque=bool(config.include_no_load_torque),
        angle_out_of_range=str(config.angle_out_of_range),
        resistance_out_of_range=str(config.resistance_out_of_range),
        max_abs_torque_Nm=float(config.max_abs_torque_Nm) if config.max_abs_torque_Nm is not None else float("inf"),
        max_electrical_power_W=float(config.max_electrical_power_W) if config.max_electrical_power_W is not None else float("inf"),
    )


def _interp_branch(branch: GeneratorMapBranch, angle_deg: float, *, out_of_range: str) -> tuple[float, float, float, bool]:
    lo = float(branch.angle_deg[0])
    hi = float(branch.angle_deg[-1])
    outside = bool(angle_deg < lo or angle_deg > hi)
    if outside and out_of_range == "error":
        raise ValueError(f"generator map angle {angle_deg:.6g} deg is outside [{lo:.6g}, {hi:.6g}] deg")
    angle = float(np.clip(angle_deg, lo, hi))
    return (
        float(np.interp(angle, branch.angle_deg, branch.torque_Nm)),
        float(np.interp(angle, branch.angle_deg, branch.voltage_V)),
        float(np.interp(angle, branch.angle_deg, branch.electrical_power_W)),
        outside,
    )


def _interp_point(point: GeneratorMapOperatingPoint, angle_deg: float, map_velocity_deg_per_s: float, policy: str) -> tuple[float, float, float, bool]:
    if map_velocity_deg_per_s > 1.0e-12:
        return _interp_branch(point.increasing, angle_deg, out_of_range=policy)
    if map_velocity_deg_per_s < -1.0e-12:
        return _interp_branch(point.decreasing, angle_deg, out_of_range=policy)
    a = _interp_branch(point.increasing, angle_deg, out_of_range=policy)
    b = _interp_branch(point.decreasing, angle_deg, out_of_range=policy)
    return (0.5 * (a[0] + b[0]), 0.5 * (a[1] + b[1]), 0.5 * (a[2] + b[2]), bool(a[3] or b[3]))


def lookup_generator_torque_map(model: GeneratorTorqueMap, q_rad: float, q_dot_rad_per_s: float, angle_min_rad: float, angle_max_rad: float) -> GeneratorMapValue:
    span = max(float(angle_max_rad) - float(angle_min_rad), 1.0e-15)
    fraction = (float(q_rad) - float(angle_min_rad)) / span
    angle_deg = model.angle_at_x_min_deg + fraction * (model.angle_at_x_max_deg - model.angle_at_x_min_deg)
    map_velocity_deg_per_s = np.degrees(float(q_dot_rad_per_s)) * (model.angle_at_x_max_deg - model.angle_at_x_min_deg) / np.degrees(span)

    resistance = float(model.load_resistance_ohm)
    r_min, r_max = float(model.resistances_ohm[0]), float(model.resistances_ohm[-1])
    resistance_outside = resistance < r_min or resistance > r_max
    if resistance_outside and model.resistance_out_of_range == "error":
        raise ValueError(f"generator load resistance {resistance:.6g} ohm is outside [{r_min:.6g}, {r_max:.6g}] ohm")
    resistance = float(np.clip(resistance, r_min, r_max))
    upper = int(np.searchsorted(model.resistances_ohm, resistance, side="right"))
    upper = min(max(upper, 1), len(model.operating_points) - 1)
    lower = upper - 1
    if len(model.operating_points) == 1:
        lower = upper = 0
    r0, r1 = float(model.resistances_ohm[lower]), float(model.resistances_ohm[upper])
    weight = 0.0 if upper == lower or r1 == r0 else (resistance - r0) / (r1 - r0)
    v0 = _interp_point(model.operating_points[lower], angle_deg, map_velocity_deg_per_s, model.angle_out_of_range)
    v1 = _interp_point(model.operating_points[upper], angle_deg, map_velocity_deg_per_s, model.angle_out_of_range)
    torque_total = (1.0 - weight) * v0[0] + weight * v1[0]
    voltage = (1.0 - weight) * v0[1] + weight * v1[1]
    electrical_power = (1.0 - weight) * v0[2] + weight * v1[2]
    no_load = _interp_point(model.no_load, angle_deg, map_velocity_deg_per_s, model.angle_out_of_range)
    load_torque = torque_total - no_load[0]
    applied_torque = torque_total if model.include_no_load_torque else load_torque
    applied_torque = float(np.clip(applied_torque, -model.max_abs_torque_Nm, model.max_abs_torque_Nm))
    electrical_power = float(np.clip(max(electrical_power, 0.0), 0.0, model.max_electrical_power_W))
    return GeneratorMapValue(
        angle_deg=float(angle_deg),
        torque_Nm=applied_torque,
        no_load_torque_Nm=float(no_load[0]),
        load_torque_Nm=float(load_torque),
        voltage_V=float(voltage),
        electrical_power_W=electrical_power,
        out_of_range=bool(resistance_outside or v0[3] or v1[3] or no_load[3]),
    )
